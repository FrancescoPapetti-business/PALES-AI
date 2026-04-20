import sys
import fitz  # PyMuPDF
import docx  # python-docx
from dataclasses import dataclass
from typing import Literal
from pathlib import Path

# Setup path come prima
BASE_DIR = Path(__file__).resolve().parents[2]
if str(BASE_DIR) not in sys.path:
    sys.path.append(str(BASE_DIR))

from core.layers.layer1_data_collection import RawDocument

# Aggiunto 'visual_guide' come richiesto per screenshot con evidenziazioni
DocType = Literal["textual", "slide", "table", "visual_guide", "mixed", "web_content"]

@dataclass
class DocumentProfile:
    raw: RawDocument
    doc_type: DocType
    quality_score: float
    has_tables: bool
    processing_strategy: str  # 'text_extract', 'vision_llm', 'ocr_tesseract'


def analyze_docx(path: str) -> dict:
    """Gestisce i file DOCX separatamente."""
    try:
        doc = docx.Document(path)
        text_len = sum(len(p.text) for p in doc.paragraphs)
        return {
            "text_len": text_len,
            "images": 0,
            "tables": len(doc.tables),
        }
    except Exception as e:
        print(f"Errore lettura DOCX {path}: {e}")
        return {"text_len": 0, "images": 0, "tables": 0}


def detect_profile(raw_doc: RawDocument) -> DocumentProfile:
    if raw_doc.source_type == "web_page":
        return DocumentProfile(
            raw=raw_doc,
            doc_type="web_content",
            quality_score=0.85,
            has_tables=False,
            processing_strategy="html_extract"
        )
    path_str = str(raw_doc.path).lower()

    # 1. Branching per tipo file
    if path_str.endswith(".docx") or path_str.endswith(".doc"):
        stats = analyze_docx(raw_doc.path)
        is_pdf = False
    else:
        # Analisi PDF
        is_pdf = True
        try:
            doc = fitz.open(raw_doc.path)
        except Exception as e:
            print(f"❌ Errore apertura PDF {raw_doc.path}: {e}")
            return DocumentProfile(
                raw=raw_doc,
                doc_type="mixed",
                quality_score=0.0,
                has_tables=False,
                processing_strategy="ocr_tesseract"
            )

        stats = {"text_len": 0, "images": 0, "tables": 0}

        for page in doc:
            text = page.get_text("text") or ""
            stats["text_len"] += len(text)
            stats["images"] += len(page.get_images())

            lines = text.split("\n")
            for line in lines:
                if line.count("|") >= 2 or (len(line) > 20 and line.count("   ") > 3):
                    stats["tables"] += 1

        doc.close()

    # 2. Classificazione aggiornata
    has_tables = stats["tables"] > 0
    score = 1.0
    strategy = "text_extract"
    doc_type = "textual"

    if is_pdf:
        # visual guide rilevata
        if stats["text_len"] < 500 and stats["images"] > 2:
            doc_type = "visual_guide"
            strategy = "vision_llm"
            score = 0.95

        elif stats["text_len"] < 500 and stats["images"] == 0:
            doc_type = "mixed"
            strategy = "ocr_tesseract"
            score = 0.5

        elif has_tables and stats["text_len"] > 500:
            doc_type = "table"
            score = 0.9

        elif stats["text_len"] < 1000 and stats["images"] > 5:
            doc_type = "slide"
            strategy = "vision_llm"
            score = 0.85

    else:
        doc_type = "textual"
        strategy = "text_extract"

    return DocumentProfile(
        raw=raw_doc,
        doc_type=doc_type,
        quality_score=max(score, 0.0),
        has_tables=has_tables,
        processing_strategy=strategy,
    )


if __name__ == "__main__":
    from core.layers.layer1_data_collection import collect_documents

    print("🔍 Testing Layer 2 (Quality Detection)...")
    docs = collect_documents()

    for d in docs[:10]:
        p = detect_profile(d)
        print(f"📄 {d.filename}")
        print(f"   → Tipo: {p.doc_type} | Strat: {p.processing_strategy} | Score: {p.quality_score}")
